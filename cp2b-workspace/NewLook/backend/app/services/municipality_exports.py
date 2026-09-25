"""Render a municipality dossier as a workbook or a PDF report.

Both renderers take what `municipality_dossier.collect()` returns and nothing
else, so the CLI export, the XLSX download and the PDF report always describe
the same municipality the same way.

Both are offered in pt-BR (the default — the dossier is a deliverable for
Brazilian planners) and in English, via ``lang``. Every word comes from
app.services.dossier_text; nothing here is copy.

The PDF's locator map is drawn from PostGIS geometry as vector paths — no tile
server, no API key, no network at render time, and it stays sharp at any zoom.
"""

from __future__ import annotations

import json
import unicodedata
from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Iterable, Sequence

import pandas as pd
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.services.dossier_text import (
    DEFAULT_LANG,
    POTENTIAL_CATEGORIES,
    RESIDUE_SECTORS,
    SECTORS,
    Lang,
    residue_name,
    sector_name,
    text,
)

SHEET_NAME_LIMIT = 31  # Excel's hard cap on sheet names.


def residues(lang: Lang = DEFAULT_LANG) -> tuple[tuple[str, str, str], ...]:
    """(key, name, sector name) for each residue, in dossier order.

    Each key is the column prefix the municipalities table uses; not every
    residue carries every column (sewage has no biogas total, forestry has no
    biomass tonnage), and a missing column reads as "—", never 0.
    """
    return tuple(
        (key, residue_name(key, lang) or key, sector_name(sector, lang) or sector)
        for key, sector in RESIDUE_SECTORS
    )


# The pt-BR tables, as the CLI exporter and older callers know them.
RESIDUOS: tuple[tuple[str, str, str], ...] = residues("pt-BR")
SETORES: tuple[tuple[str, str], ...] = tuple((k, sector_name(k, "pt-BR") or k) for k in SECTORS)

HEADER_FILL = "FF1F7A4D"
HEADER_FONT = "FFFFFFFF"
BAND_FILL = "FFF2F7F4"

BRAND_GREEN = colors.HexColor("#1f7a4d")
BRAND_LIME = colors.HexColor("#8bc34a")
INK = colors.HexColor("#1a1a1a")
MUTED = colors.HexColor("#6b7280")
HAIRLINE = colors.HexColor("#d1d5db")

# The locator map box, in points.
MAP_W, MAP_H = 165 * mm, 110 * mm


def ascii_slug(name: str) -> str:
    """ "Lençóis Paulista" -> "lencois_paulista", for filenames and headers.

    Content-Disposition is ASCII by default, so an accented filename either gets
    mangled or needs RFC 5987 encoding that older clients mishandle.
    """
    decomposed = unicodedata.normalize("NFKD", str(name))
    stripped = "".join(c for c in decomposed if not unicodedata.combining(c))
    return "".join(c if c.isalnum() else "_" for c in stripped.lower()).strip("_")


def _naive(value: Any) -> Any:
    """Excel has no timezone-aware datetime and openpyxl refuses one outright."""
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    return value


def _num(value: Any) -> float | None:
    """A number, or None. Never 0 as a stand-in for absent — the two differ."""
    if value is None:
        return None
    try:
        out = float(value)
    except (TypeError, ValueError):
        return None
    return None if pd.isna(out) else out


def potential_category(code: Any, lang: Lang = DEFAULT_LANG) -> Any:
    """The backend's ALTO / MEDIO / BAIXO code, worded for the reader."""
    if code is None:
        return None
    return POTENTIAL_CATEGORIES.get(lang, POTENTIAL_CATEGORIES[DEFAULT_LANG]).get(str(code), code)


def _sheet_resumo(
    who: dict, muni: dict, lang: Lang = DEFAULT_LANG, cp2b: dict | None = None
) -> pd.DataFrame:
    """Identification and the headline figures, the one page most readers need.

    `cp2b` is the municipality's row of the CP2b view: Cenário Real is N4 and
    Cenário Ideal N3. Absent (outside SP, or before migration 034) it leaves the
    two cells empty rather than zero.
    """
    t = text(lang)
    cp2b = cp2b or {}
    rows = [
        (t["sec_identification"], None, None),
        (t["municipality"], who.get("municipality_name"), None),
        (t["ibge_code"], who.get("ibge_code"), None),
        (t["uf"], who.get("uf"), None),
        (t["immediate_region"], who.get("immediate_region"), None),
        (t["intermediate_region"], who.get("intermediate_region"), None),
        (t["area"], _num(who.get("area_km2")), "km²"),
        (
            t["population"],
            _num(who.get("population")),
            t["population_unit"].format(year=who.get("population_year") or "—"),
        ),
        (t["density"], _num(muni.get("population_density")), t["density_unit"]),
        (None, None, None),
        (t["sec_biogas"], None, None),
        (t["note"], t["legacy_value"], None),
        (t["total_potential"], _num(muni.get("total_biogas_m3_year")), t["unit_m3_year"]),
        (t["daily_potential"], _num(muni.get("total_biogas_m3_day")), t["unit_m3_day"]),
        (t["potential_class"], potential_category(who.get("potential_category"), lang), None),
        (None, None, None),
        (t["sec_energy"], None, None),
        (t["energy_potential"], _num(muni.get("energy_potential_mwh_year")), t["unit_mwh_year"]),
        (t["energy_potential"], _num(muni.get("energy_potential_kwh_day")), t["unit_kwh_day"]),
        (t["co2_avoided"], _num(muni.get("co2_reduction_tons_year")), t["unit_t_year"]),
        (None, None, None),
        (t["sec_biomass"], None, None),
        (t["total_biomass"], _num(muni.get("total_biomass_tons_year")), t["unit_t_year"]),
        (None, None, None),
        (t["sec_methane"], None, None),
        (t["ch4_real"], _num(cp2b.get("ch4_cp2b_n4_m3_year")), t["unit_m3_year"]),
        (t["ch4_ideal"], _num(cp2b.get("ch4_cp2b_n3_m3_year")), t["unit_m3_year"]),
        (t["note"], t["cp2b_value"], None),
        (None, None, None),
        (t["sec_provenance"], None, None),
        (t["data_confidence"], who.get("data_confidence"), None),
        (t["extracted_at"], datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M"), None),
        (t["source"], "PILAR-2b", None),
    ]
    return pd.DataFrame(rows, columns=[t["col_indicator"], t["col_value"], t["col_unit"]])


def _sheet_setores(muni: dict, lang: Lang = DEFAULT_LANG, cp2b: dict | None = None) -> pd.DataFrame:
    """Biogas, biomass and both CH4 scenarios per sector, with each one's share.

    The scenarios are CP2b (Real = N4, Ideal = N3); the method has no forestry
    stream, so that row's scenario cells stay empty.
    """
    t = text(lang)
    cp2b = cp2b or {}
    biomass, biogas = t["col_biomass"], t["col_biogas"]
    real, ideal, share = t["col_ch4_real"], t["col_ch4_ideal"], t["col_share"]
    rows = []
    for key in SECTORS:
        rows.append(
            {
                t["col_sector"]: sector_name(key, lang) or key,
                biomass: _num(muni.get(f"{key}_biomass_tons_year")),
                biogas: _num(muni.get(f"{key}_biogas_m3_year")),
                real: _num(cp2b.get(f"ch4_cp2b_n4_{key}_m3_year")),
                ideal: _num(cp2b.get(f"ch4_cp2b_n3_{key}_m3_year")),
            }
        )
    frame = pd.DataFrame(rows)
    # min_count=1: a column with no value at all totals to empty, not 0. Without
    # it a municipality with no CP2b row (outside SP, or before migration 034 is
    # loaded) got a TOTAL of 0 under Real/Ideal while every sector cell was blank.
    total = frame[biogas].sum(skipna=True, min_count=1)
    has_total = pd.notna(total) and total != 0
    frame[share] = frame[biogas] / total if has_total else None
    total_row = {
        t["col_sector"]: t["total"],
        biomass: frame[biomass].sum(skipna=True, min_count=1),
        biogas: total,
        real: frame[real].sum(skipna=True, min_count=1),
        ideal: frame[ideal].sum(skipna=True, min_count=1),
        share: 1.0 if has_total else None,
    }
    return pd.concat([frame, pd.DataFrame([total_row])], ignore_index=True)


def _coalesce(*values: Any) -> float | None:
    """First value that is actually present.

    Deliberately not `a or b`: a measured zero is falsy, and 0 t/year of
    aquaculture residue is a finding, not a missing reading. Conflating the two
    is the bug this function exists to prevent.
    """
    for value in values:
        number = _num(value)
        if number is not None:
            return number
    return None


def _sheet_residuos(
    muni: dict, streams: list[dict], lang: Lang = DEFAULT_LANG, cp2b: dict | None = None
) -> pd.DataFrame:
    """One row per residue, joining the municipality columns to the SP stream table.

    The stream table carries energy and the conversion factor actually applied;
    the CP2b view carries the CH4 scenarios (Real = N4, Ideal = N3). Neither has
    both, so the sheet is the join — which is the whole reason this export exists.
    """
    t = text(lang)
    cp2b = cp2b or {}
    biogas = t["col_biogas"]
    by_stream = {str(r.get("residue_stream") or "").lower(): r for r in streams}
    rows = []
    for key, label, setor in residues(lang):
        stream = by_stream.get(key, {})
        rows.append(
            {
                t["col_residue"]: label,
                t["col_sector"]: setor,
                t["col_biomass"]: _coalesce(
                    muni.get(f"{key}_biomass_tons_year"), stream.get("residue_tons_yr")
                ),
                biogas: _coalesce(muni.get(f"{key}_biogas_m3_year"), stream.get("biogas_m3_yr")),
                t["col_ch4_real"]: _num(cp2b.get(f"ch4_cp2b_n4_{key}_m3_year")),
                t["col_ch4_ideal"]: _num(cp2b.get(f"ch4_cp2b_n3_{key}_m3_year")),
                t["col_energy"]: _num(stream.get("energy_mwh_yr")),
                t["col_factor"]: _num(stream.get("conversion_factor")),
                t["col_factor_unit"]: stream.get("cf_unit"),
            }
        )
    frame = pd.DataFrame(rows)
    total = frame[biogas].sum(skipna=True)
    frame[t["col_share"]] = frame[biogas] / total if total else None
    return frame.sort_values(biogas, ascending=False, na_position="last")


def _sheet_series(series: list[dict], lang: Lang = DEFAULT_LANG) -> pd.DataFrame:
    """The time series, tidied and sorted — year, variable, value, unit, source."""
    t = text(lang)
    if not series:
        return pd.DataFrame({t["no_series"]: []})
    frame = pd.DataFrame(series)
    keep = {
        "year": t["col_year"],
        "variable": t["col_variable"],
        "value": t["col_value"],
        "unit": t["col_unit"],
        "quality": t["col_quality"],
        "source_id": t["source"],
    }
    frame = frame[[c for c in keep if c in frame.columns]].rename(columns=keep)
    order = [c for c in (t["source"], t["col_variable"], t["col_year"]) if c in frame.columns]
    return frame.sort_values(order, na_position="last") if order else frame


def _sheet_infra(features: list[dict], lang: Lang = DEFAULT_LANG) -> pd.DataFrame:
    t = text(lang)
    if not features:
        return pd.DataFrame({t["no_infra"]: []})
    frame = pd.DataFrame(features)
    keep = {
        "layer_id": t["col_layer"],
        "name": t["col_name"],
        "source_id": t["source"],
        "attributes": t["col_attributes"],
    }
    frame = frame[[c for c in keep if c in frame.columns]].rename(columns=keep)
    attributes = t["col_attributes"]
    if attributes in frame.columns:
        frame[attributes] = frame[attributes].map(
            lambda v: json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v
        )
    return frame


def _sheet_fontes(
    sections: dict[str, list[dict]], who: dict, lang: Lang = DEFAULT_LANG
) -> pd.DataFrame:
    """Where the numbers came from, and what is missing — stated, not implied."""
    t = text(lang)
    rows: list[tuple[str, Any]] = [(t["sec_data_provenance"], None)]
    for record in sections.get("biomass_provenance") or []:
        label = record.get("residue") or record.get("field") or record.get("column_name") or "—"
        rows.append((str(label), record.get("source") or record.get("source_id") or "—"))

    sources = sorted(
        {str(r.get("source_id")) for r in sections.get("timeseries") or [] if r.get("source_id")}
    )
    if sources:
        rows += [(None, None), (t["sec_series_sources"], None)]
        rows += [(s, None) for s in sources]

    typology = (sections.get("typology") or [{}])[0]
    rows += [(None, None), (t["sec_typology"], None)]
    rows.append((t["cn_molar"], typology.get("cn_molar")))
    rows.append((t["cn_harm"], typology.get("cn_harm")))
    if typology.get("cn_harm") is None:
        # Say it outright. A blank cell invites the reader to assume zero.
        rows.append((t["note"], t["cn_harm_missing"]))

    rows += [
        (None, None),
        (t["sec_notes"], None),
        (t["nature"], t["nature_value"]),
        (t["cp2b_label"], t["cp2b_value"]),
        (t["legacy_label"], t["legacy_value"]),
        (t["data_confidence"], who.get("data_confidence")),
    ]
    return pd.DataFrame(rows, columns=[t["col_item"], t["col_detail"]])


def _style(
    writer,
    sheet_name: str,
    frame: pd.DataFrame,
    widths: Sequence[int] | None = None,
    decimal_columns: frozenset[str] = frozenset(),
) -> None:
    """Header band, frozen top row, sensible widths and thousands separators.

    ``decimal_columns`` keep two decimals (conversion factors); every other
    numeric column is written as a whole number, every "%" column as a share.
    """
    from openpyxl.styles import Alignment, Font, PatternFill

    ws = writer.sheets[sheet_name]
    fill = PatternFill("solid", fgColor=HEADER_FILL)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(bold=True, color=HEADER_FONT)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28

    for index, column in enumerate(frame.columns, start=1):
        letter = ws.cell(row=1, column=index).column_letter
        if widths and index <= len(widths):
            ws.column_dimensions[letter].width = widths[index - 1]
        else:
            longest = max([len(str(column))] + [len(str(v)) for v in frame[column].head(60)])
            ws.column_dimensions[letter].width = min(max(longest + 2, 12), 46)

        if "%" in str(column):
            number_format = "0.0%"
        elif frame[column].dtype.kind in "fi":
            number_format = "#,##0.00" if column in decimal_columns else "#,##0"
        else:
            continue
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=index).number_format = number_format

    # Bold the section captions and the TOTAL line — they orient the reader.
    for row in range(2, ws.max_row + 1):
        first = ws.cell(row=row, column=1).value
        if isinstance(first, str) and (first.isupper() and len(first) > 3):
            ws.cell(row=row, column=1).font = Font(bold=True, color="FF1F7A4D")


def build_workbook(sections: dict[str, list[dict]], who: dict, lang: Lang = DEFAULT_LANG) -> bytes:
    """The dossier as a curated workbook: six sheets, no raw table dumps.

    Deliberately not one sheet per database table. The municipalities row alone is
    98 columns wide and means nothing to a reader; the value is in reshaping it
    into sector and residue tables and joining the SP stream data onto it.
    """
    t = text(lang)
    muni = (sections.get("municipality") or [{}])[0]
    cp2b = (sections.get("cp2b") or [{}])[0]
    streams = sections.get("residue_streams") or []
    sheets: list[tuple[str, pd.DataFrame, Sequence[int] | None]] = [
        (t["sheet_summary"], _sheet_resumo(who, muni, lang, cp2b), (34, 20, 22)),
        (t["sheet_sectors"], _sheet_setores(muni, lang, cp2b), (16, 20, 20, 20, 20, 14)),
        (t["sheet_residues"], _sheet_residuos(muni, streams, lang, cp2b), None),
        (
            t["sheet_series"],
            _sheet_series(sections.get("timeseries") or [], lang),
            (10, 30, 16, 14, 14, 18),
        ),
        (
            t["sheet_infra"],
            _sheet_infra(sections.get("infrastructure") or [], lang),
            (22, 34, 16, 50),
        ),
        (t["sheet_sources"], _sheet_fontes(sections, who, lang), (40, 62)),
    ]
    decimals = frozenset({t["col_factor"]})

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame, widths in sheets:
            safe = frame.map(_naive)
            safe.to_excel(writer, sheet_name=name[:SHEET_NAME_LIMIT], index=False)
            _style(writer, name[:SHEET_NAME_LIMIT], safe, widths, decimals)
    return buffer.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────────────


def _styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "title",
            parent=base["Title"],
            fontSize=22,
            leading=26,
            textColor=BRAND_GREEN,
            alignment=0,
            spaceAfter=2,
        ),
        "subtitle": ParagraphStyle(
            "subtitle",
            parent=base["Normal"],
            fontSize=10.5,
            leading=14,
            textColor=MUTED,
        ),
        "h2": ParagraphStyle(
            "h2",
            parent=base["Heading2"],
            fontSize=13,
            leading=16,
            textColor=BRAND_GREEN,
            spaceBefore=12,
            spaceAfter=6,
        ),
        "body": ParagraphStyle(
            "body", parent=base["Normal"], fontSize=9.5, leading=13, textColor=INK
        ),
        "note": ParagraphStyle(
            "note", parent=base["Normal"], fontSize=8, leading=11, textColor=MUTED
        ),
    }


def _ring(geometry: dict | None) -> list[list[Sequence[float]]]:
    """Flatten a GeoJSON Polygon/MultiPolygon into a list of exterior rings."""
    if not geometry:
        return []
    kind, coords = geometry.get("type"), geometry.get("coordinates") or []
    if kind == "Polygon":
        return [coords[0]] if coords else []
    if kind == "MultiPolygon":
        return [poly[0] for poly in coords if poly]
    return []


def _fit(rings: Iterable[Sequence[Sequence[float]]], width: float, height: float):
    """Return a lon/lat -> canvas projector that fits `rings` into the box.

    A plain equirectangular fit with one shared scale, so the state does not come
    out stretched. São Paulo at this size does not warrant a real projection.
    """
    xs = [p[0] for r in rings for p in r]
    ys = [p[1] for r in rings for p in r]
    if not xs or not ys:
        return None
    min_x, max_x, min_y, max_y = min(xs), max(xs), min(ys), max(ys)
    span_x = (max_x - min_x) or 1e-9
    span_y = (max_y - min_y) or 1e-9
    scale = min(width / span_x, height / span_y) * 0.92
    off_x = (width - span_x * scale) / 2
    off_y = (height - span_y * scale) / 2

    def project(lon: float, lat: float) -> tuple[float, float]:
        return (off_x + (lon - min_x) * scale, off_y + (lat - min_y) * scale)

    return project


def _locator_map(
    state_geojson: str | None, muni_geojson: str | None, uf: str, lang: Lang = DEFAULT_LANG
):
    """The municipality picked out inside its state outline, as vector paths."""
    from reportlab.graphics.shapes import Drawing, Polygon, String

    drawing = Drawing(MAP_W, MAP_H)
    state = json.loads(state_geojson) if state_geojson else None
    muni = json.loads(muni_geojson) if muni_geojson else None

    state_rings = _ring(state)
    muni_rings = _ring(muni)
    project = _fit(state_rings or muni_rings, MAP_W, MAP_H)
    if project is None:
        drawing.add(
            String(
                8, MAP_H / 2, text(lang)["pdf_geometry_unavailable"], fontSize=9, fillColor=MUTED
            )
        )
        return drawing

    for ring in state_rings:
        drawing.add(
            Polygon(
                [c for p in ring for c in project(p[0], p[1])],
                fillColor=colors.HexColor("#eef2f0"),
                strokeColor=HAIRLINE,
                strokeWidth=0.6,
            )
        )
    for ring in muni_rings:
        drawing.add(
            Polygon(
                [c for p in ring for c in project(p[0], p[1])],
                fillColor=BRAND_LIME,
                strokeColor=BRAND_GREEN,
                strokeWidth=1.1,
            )
        )
    caption = text(lang)["pdf_location_within"].format(uf=uf)
    drawing.add(String(4, 4, caption, fontSize=7.5, fillColor=MUTED))
    return drawing


def _fmt(value: Any, digits: int = 0, lang: Lang = DEFAULT_LANG) -> str:
    """Space-grouped thousands (read the same in both languages; a plain space,
    since the PDF's base fonts have no thin space) and the language's decimal
    mark: 1 234.5 (en) / 1 234,5 (pt-BR)."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return "—"
    if isinstance(value, (int, float)):
        grouped = f"{float(value):,.{digits}f}"
        decimal = "," if lang == "pt-BR" else "."
        return grouped.replace(",", " ").replace(".", decimal)
    return str(value)


def _kv_table(rows: list[tuple[str, str]], widths=(58 * mm, 42 * mm)) -> Table:
    table = Table(rows, colWidths=list(widths), hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("TEXTCOLOR", (0, 0), (0, -1), MUTED),
                ("TEXTCOLOR", (1, 0), (1, -1), INK),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("LINEBELOW", (0, 0), (-1, -2), 0.25, HAIRLINE),
            ]
        )
    )
    return table


def _data_table(header: Sequence[str], body: Sequence[Sequence[Any]], widths) -> Table:
    table = Table([list(header)] + [list(r) for r in body], colWidths=list(widths), hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), BRAND_GREEN),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7faf8")]),
                ("GRID", (0, 0), (-1, -1), 0.25, HAIRLINE),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def _stream_label(stream: dict, lang: Lang) -> str:
    """A residue stream by its key, in ``lang``; the served names are the fallback."""
    key = str(stream.get("residue_stream") or "").lower()
    return residue_name(key, lang) or str(
        stream.get("residue_stream_pt") or stream.get("residue_stream") or "—"
    )


def _stream_sector(stream: dict, lang: Lang) -> str:
    key = str(stream.get("sector") or "").lower()
    return sector_name(key, lang) or str(stream.get("sector_pt") or stream.get("sector") or "—")


def build_pdf(
    sections: dict[str, list[dict]],
    who: dict,
    muni_geojson: str | None = None,
    state_geojson: str | None = None,
    lang: Lang = DEFAULT_LANG,
) -> bytes:
    """A two-page municipal report: headline figures, locator map, residue detail."""
    t = text(lang)

    def fmt(value: Any, digits: int = 0) -> str:
        return _fmt(value, digits, lang)

    style = _styles()
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"PILAR-2b — {who.get('municipality_name')}",
        author="PILAR-2b",
    )

    muni = (sections.get("municipality") or [{}])[0]
    cp2b = (sections.get("cp2b") or [{}])[0]
    story: list[Any] = []

    story.append(Paragraph(str(who.get("municipality_name", "—")), style["title"]))
    story.append(
        Paragraph(
            f"{who.get('uf')} · IBGE {who.get('ibge_code')} · "
            f"{who.get('immediate_region')} / {who.get('intermediate_region')}",
            style["subtitle"],
        )
    )
    story.append(Spacer(1, 10))

    population_year = who.get("population_year")
    headline = [
        (t["ch4_real"], f"{fmt(cp2b.get('ch4_cp2b_n4_m3_year'))} {t['unit_m3_year']}"),
        (t["ch4_ideal"], f"{fmt(cp2b.get('ch4_cp2b_n3_m3_year'))} {t['unit_m3_year']}"),
        (t["pdf_total_biogas"], f"{fmt(muni.get('total_biogas_m3_year'))} {t['unit_m3_year']}"),
        (t["pdf_energy"], f"{fmt(muni.get('energy_potential_mwh_year'))} {t['unit_mwh_year']}"),
        (t["pdf_co2"], f"{fmt(muni.get('co2_reduction_tons_year'))} {t['unit_t_year']}"),
        (t["pdf_biomass"], f"{fmt(muni.get('total_biomass_tons_year'))} {t['unit_t_year']}"),
        (t["pdf_category"], str(potential_category(who.get("potential_category"), lang) or "—")),
        # A year is a label, not a quantity: no thousands separator.
        (t["pdf_population"], f"{fmt(who.get('population'))} ({population_year or '—'})"),
        (t["pdf_area"], f"{fmt(who.get('area_km2'), 1)} km²"),
        (t["pdf_confidence"], str(who.get("data_confidence") or "—")),
    ]
    story.append(_kv_table(headline))
    story.append(Paragraph(t["pdf_method_note"], style["note"]))
    story.append(Spacer(1, 12))

    story.append(Paragraph(t["pdf_location"], style["h2"]))
    story.append(_locator_map(state_geojson, muni_geojson, str(who.get("uf") or ""), lang))

    # ── Sector breakdown ─────────────────────────────────────────────────────
    sectors = [(sector_name(k, lang) or k, muni.get(f"{k}_biogas_m3_year")) for k in SECTORS]
    total = sum(float(v or 0) for _, v in sectors) or 1.0
    story.append(PageBreak())
    story.append(Paragraph(t["pdf_by_sector"], style["h2"]))
    story.append(
        _data_table(
            [t["pdf_col_sector"], t["pdf_col_m3_year"], t["pdf_col_share"]],
            [(n, fmt(v), f"{fmt(100 * float(v or 0) / total, 1)}%") for n, v in sectors],
            widths=(70 * mm, 55 * mm, 25 * mm),
        )
    )
    story.append(Spacer(1, 12))

    # ── Residue streams ──────────────────────────────────────────────────────
    streams = sections.get("residue_streams") or []
    if streams:
        ordered = sorted(streams, key=lambda r: float(r.get("biogas_m3_yr") or 0), reverse=True)
        story.append(Paragraph(t["pdf_streams"], style["h2"]))
        story.append(
            _data_table(
                [
                    t["pdf_col_stream"],
                    t["pdf_col_sector"],
                    t["pdf_col_t_year"],
                    t["pdf_col_biogas"],
                    t["pdf_col_mwh"],
                ],
                [
                    (
                        _stream_label(r, lang),
                        _stream_sector(r, lang),
                        fmt(r.get("residue_tons_yr")),
                        fmt(r.get("biogas_m3_yr")),
                        fmt(r.get("energy_mwh_yr")),
                    )
                    for r in ordered
                ],
                widths=(38 * mm, 30 * mm, 30 * mm, 38 * mm, 28 * mm),
            )
        )
        story.append(Spacer(1, 10))

    # ── Provenance ───────────────────────────────────────────────────────────
    # A report that cannot say where its numbers came from is not much use in a
    # planning meeting, which is where these are meant to be read.
    provenance = sections.get("biomass_provenance") or []
    series = sections.get("timeseries") or []
    sources = sorted({str(r.get("source_id")) for r in series if r.get("source_id")})
    if provenance or sources:
        bits = []
        if sources:
            bits.append(t["pdf_series_sources"].format(sources=", ".join(sources)))
        if provenance:
            bits.append(t["pdf_provenance_count"].format(count=len(provenance)))
        story.append(
            KeepTogether(
                [
                    Paragraph(t["pdf_provenance"], style["h2"]),
                    Paragraph(" ".join(bits), style["body"]),
                ]
            )
        )

    story.append(Spacer(1, 14))
    when = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    story.append(Paragraph(t["pdf_generated"].format(when=when), style["note"]))

    doc.build(story)
    return buffer.getvalue()
